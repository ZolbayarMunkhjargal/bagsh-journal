"""
Validates the .docx and .xlsx files KinderJournal writes.

The application produces these formats without a document library, so the only
convincing check is to open the files with a real one. python-docx and openpyxl
are independent implementations of the same formats, so if they accept the files
and read the expected content back, the structure is sound.

    python tools/check_office_files.py <folder>
"""

from __future__ import annotations

import html
import re
import sys
import zipfile
from pathlib import Path

import docx
import openpyxl

failures: list[str] = []
checked = 0


def fail(message: str) -> None:
    failures.append(message)
    print(f"  FAIL  {message}")


def ok(message: str) -> None:
    print(f"  ok    {message}")


def check_zip(path: Path, required: list[str]) -> None:
    """A missing part is what makes Word or Excel call a file corrupt."""
    try:
        with zipfile.ZipFile(path) as archive:
            names = archive.namelist()
            bad = archive.testzip()
    except zipfile.BadZipFile as exc:
        fail(f"{path.name}: not a zip archive ({exc})")
        return
    if bad:
        fail(f"{path.name}: corrupt entry {bad}")
        return
    missing = [name for name in required if name not in names]
    if missing:
        fail(f"{path.name}: missing part(s) {missing}")
    else:
        ok(f"{path.name}: valid archive with {len(names)} parts")


def check_docx(path: Path) -> None:
    check_zip(path, ["[Content_Types].xml", "_rels/.rels", "word/document.xml",
                     "word/styles.xml"])
    try:
        document = docx.Document(str(path))
    except Exception as exc:                       # noqa: BLE001 - report anything
        fail(f"{path.name}: python-docx could not open it ({exc})")
        return
    text = "\n".join(p.text for p in document.paragraphs if p.text.strip())
    tables = document.tables
    total_cells = sum(len(t.rows) * len(t.columns) for t in tables)
    if not text.strip():
        fail(f"{path.name}: opened but has no text")
        return
    if "{" in text and "}" in text:
        fail(f"{path.name}: text contains an unreplaced placeholder")
    ok(f"{path.name}: {len(text.splitlines())} paragraphs, {len(tables)} tables, "
       f"{total_cells} cells")
    # A heading style is what makes Word's navigation pane work.
    styles = {p.style.name.lower() for p in document.paragraphs if p.text.strip()}
    if any("heading" in s for s in styles):
        ok(f"{path.name}: uses heading styles ({sorted(styles)[:3]})")
    else:
        fail(f"{path.name}: no heading styles found")


def check_xlsx(path: Path) -> None:
    check_zip(path, ["[Content_Types].xml", "_rels/.rels", "xl/workbook.xml",
                     "xl/styles.xml"])
    try:
        book = openpyxl.load_workbook(str(path))
    except Exception as exc:                       # noqa: BLE001 - report anything
        fail(f"{path.name}: openpyxl could not open it ({exc})")
        return
    if not book.sheetnames:
        fail(f"{path.name}: no sheets")
        return
    rows = 0
    numeric = 0
    text_cells = 0
    for name in book.sheetnames:
        sheet = book[name]
        rows += sheet.max_row
        for row in sheet.iter_rows(values_only=True):
            for value in row:
                if isinstance(value, (int, float)) and not isinstance(value, bool):
                    numeric += 1
                elif isinstance(value, str):
                    text_cells += 1
    ok(f"{path.name}: {len(book.sheetnames)} sheet(s) {book.sheetnames[:3]}, "
       f"{rows} rows, {numeric} numeric cells")
    if rows < 2:
        fail(f"{path.name}: no data rows")
    # Numbers must be stored as numbers so Excel can total them, but a report can
    # legitimately have none - a search result listing dates and names, for
    # instance - so this is only a problem when a table held a real quantity.
    # Only the tables are examined: a caption such as "returned 96 records" is a
    # sentence and belongs in the document, not in a cell.
    source = path.with_suffix(".html")
    if numeric == 0 and source.exists():
        raw = source.read_text(encoding="utf-8", errors="replace")
        cells = " ".join(
            re.sub(r"<[^>]+>", " ", table)
            for table in re.findall(r"<table.*?</table>", raw, re.S)
        )
        remaining = without_dates(cells)
        if re.search(r"\d", remaining):
            fail(f"{path.name}: numbers in the report are stored as text, "
                 f"so Excel could not total them")
        else:
            ok(f"{path.name}: this report holds no quantities, so nothing to total")


def without_dates(text: str) -> str:
    """Removes dates and clock times, leaving only what could be a quantity."""
    # Entities first: an escaped apostrophe is "&#39;", whose digits would
    # otherwise look like a number the spreadsheet had failed to store as one.
    text = html.unescape(text)
    months = "Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec"
    # "12 Oct 2026" and "October 2026".
    text = re.sub(rf"\b\d{{1,2}}\s+(?:{months})[a-z]*\.?\s+\d{{4}}\b", " ", text)
    text = re.sub(rf"\b(?:{months})[a-z]*\.?\s+\d{{4}}\b", " ", text)
    # Clock times, ISO dates, and a year standing on its own as a label.
    text = re.sub(r"\b\d{1,2}:\d{2}\b", " ", text)
    text = re.sub(r"\b\d{4}-\d{2}-\d{2}\b", " ", text)
    text = re.sub(r"\b(?:19|20)\d{2}\b", " ", text)
    return text


def main() -> None:
    folder = Path(sys.argv[1] if len(sys.argv) > 1 else "build/office-check")
    docx_files = sorted(folder.glob("*.docx"))
    xlsx_files = sorted(folder.glob("*.xlsx"))
    if not docx_files and not xlsx_files:
        print(f"No .docx or .xlsx files in {folder}")
        raise SystemExit(1)

    print(f"Checking {len(docx_files)} Word and {len(xlsx_files)} Excel files "
          f"in {folder}")
    for path in docx_files:
        check_docx(path)
    for path in xlsx_files:
        check_xlsx(path)

    print()
    if failures:
        print(f"{len(failures)} problem(s):")
        for message in failures:
            print(f"  - {message}")
        raise SystemExit(1)
    print("All files opened and read correctly.")


if __name__ == "__main__":
    main()
